"""
IAR Excel placeholder filling utilities.

This module is a lightweight migration of the Excel-related placeholder
replacement logic from datamerge/util.py. It is intended for the IAR workflow:
QSCL0415_Installation_Assessment_Review_v3.2.1.xlsx and similar Excel templates.

Supported placeholders:
    <PMS.customer>
    <PMS.project>
    <PMS.sop>
    <PMS.PJM>
    <PMS.TPM>
    <PMS.SW_PCM>
    <PMS.customer-project>

Rules:
    1. Only Excel files are handled here.
    2. The whole workbook is scanned: all sheets, all cells.
    3. Unknown placeholders are replaced with "N/A".
    4. Combined placeholders such as <PMS.customer-project> are joined by "_".
    5. .xlsm files try to preserve VBA content.
"""

from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path
from typing import Any, Union

import openpyxl
from fastapi import HTTPException


_PLACEHOLDER_PATTERN = re.compile(r"<PMS\.([\w\-]+)>")


_ROLE_MAPPING = {
    "PJM": "PjM",
    "TPM": "TPM",
    "ECU_PCM": "ECU-PCM",
    "SW_PCM": "SW_PCM",
    "FSM": "FSM",
    "SYS_ENG": "Sys-ENG",
    "APP_PCM": "App PCM",
    "HW_Dev": "HW Developer",
    "AM": "AM",
    "CM": "CM",
    "COS": "COS",
    "MECH_PCM": "MECH-PCM",
    "SAMCO": "SAMCO",
    "SEC": "SEC",
    "TestM": "Test Manager",
}


def _parse_role_email_summary(role_email_summary: str) -> dict[str, str]:
    """
    Parse role_email_summary into fields usable by <PMS.xxxEmail>.

    Expected examples:
        PjMEmail: Zhang San_zhang.san@example.com;
        TPMEmail: Li Si_li.si@example.com

    Output example:
        {
            "PjMEmail": "Zhang San\nzhang.san@example.com",
            "TPMEmail": "Li Si\nli.si@example.com",
        }
    """
    if not role_email_summary or role_email_summary == "N/A":
        return {}

    email_data_by_role: dict[str, str] = {}

    for section in str(role_email_summary).split(";"):
        if ":" not in section:
            continue

        role_key, _, member_list_str = section.partition(":")
        role_key = role_key.strip()
        if not role_key:
            continue

        output_lines: list[str] = []
        for member_str in member_list_str.strip().split(","):
            member_str = member_str.strip()
            if not member_str:
                continue

            display_name, separator, email = member_str.rpartition("_")
            if separator:
                output_lines.append(display_name.strip())
                output_lines.append(email.strip())
            else:
                output_lines.append(member_str)

        if output_lines:
            email_data_by_role[role_key] = "\n".join(output_lines)

    return email_data_by_role


def _prepare_profile_for_filling(profile_dict: dict[str, Any]) -> dict[str, str]:
    """
    Normalize project data and expand role placeholders.

    It converts the original project profile into a flat placeholder dictionary.
    For example:
        role_summary = "PjM: Zhang San; TPM: Li Si"
    becomes:
        PJM = "Zhang San"
        TPM = "Li Si"
    """
    formatted_profile: dict[str, str] = {}

    for key, value in (profile_dict or {}).items():
        if value is None or value == "" or value == []:
            formatted_profile[str(key)] = "N/A"
        elif isinstance(value, (list, tuple, set)):
            formatted_profile[str(key)] = ", ".join(map(str, value))
        else:
            formatted_profile[str(key)] = str(value)

    # Parse role_summary for role names.
    role_summary_str = formatted_profile.get("role_summary", "N/A")
    role_data: dict[str, str] = {}
    if role_summary_str and role_summary_str != "N/A":
        for pair in role_summary_str.split(";"):
            if ":" in pair:
                role, _, names = pair.partition(":")
                role_data[role.strip()] = names.strip() or "N/A"

    for placeholder_key, summary_key in _ROLE_MAPPING.items():
        formatted_profile[placeholder_key] = role_data.get(summary_key, "N/A")

    # Parse role_email_summary for role emails.
    inverted_role_mapping = {value: key for key, value in _ROLE_MAPPING.items()}
    email_summary_str = formatted_profile.get("role_email_summary", "N/A")
    parsed_emails_by_role = _parse_role_email_summary(email_summary_str)

    corrected_email_data: dict[str, str] = {}
    for raw_key, value in parsed_emails_by_role.items():
        if raw_key.endswith("Email"):
            prefix = raw_key[:-5]
            placeholder_prefix = inverted_role_mapping.get(prefix)
            if placeholder_prefix:
                corrected_email_data[placeholder_prefix + "Email"] = value
            else:
                corrected_email_data[raw_key] = value
        else:
            corrected_email_data[raw_key] = value

    formatted_profile.update(corrected_email_data)
    return formatted_profile


def _is_macro_enabled(source: Union[Path, io.BytesIO]) -> bool:
    """Return True for .xlsm files or in-memory Excel files containing VBA."""
    if isinstance(source, Path):
        return source.suffix.lower() == ".xlsm"

    if isinstance(source, io.BytesIO):
        try:
            source.seek(0)
            with zipfile.ZipFile(source, "r") as zip_file:
                return "xl/vbaProject.bin" in zip_file.namelist()
        except zipfile.BadZipFile:
            return False
        finally:
            source.seek(0)

    return False


def _replace_placeholder(match: re.Match[str], formatted_profile: dict[str, str]) -> str:
    combined_key = match.group(1)

    if "-" in combined_key:
        keys = combined_key.split("-")
        value_parts = [formatted_profile.get(key, "N/A") for key in keys]
        return "_".join(value_parts)

    return formatted_profile.get(combined_key, "N/A")


def fill_excel_by_placeholders(
    profile_dict: dict[str, Any],
    source: Union[Path, io.BytesIO],
) -> io.BytesIO:
    """
    Fill all <PMS.xxx> placeholders in an Excel workbook.

    Parameters:
        profile_dict:
            Project data dictionary used for placeholder replacement.
        source:
            Path or BytesIO of the Excel template.

    Returns:
        BytesIO of the filled Excel workbook.
    """
    should_keep_vba = _is_macro_enabled(source)

    try:
        if isinstance(source, io.BytesIO):
            source.seek(0)
        workbook = openpyxl.load_workbook(source, keep_vba=should_keep_vba)
    except FileNotFoundError as exc:
        raise HTTPException(
            status_code=500,
            detail=f"IAR template file not found: {source}",
        ) from exc
    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail=f"无法处理提供的 IAR Excel 文件: {exc}",
        ) from exc

    formatted_profile = _prepare_profile_for_filling(profile_dict)

    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and "<PMS." in cell.value:
                    cell.value = _PLACEHOLDER_PATTERN.sub(
                        lambda match: _replace_placeholder(match, formatted_profile),
                        cell.value,
                    )

    stream = io.BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream


def build_demo_iar_profile() -> dict[str, Any]:
    """Small local test profile. This function is only for manual testing."""
    return {
        "customer": "Geely",
        "project": "BX11",
        "model": "BX11 HEV",
        "sop": "2027-08-01",
        "plattform": "PPE",
        "project_leader": "Zhang San",
        "vint_responsible": "Li Si",
        "peripheral_sensor_configuration": "UFS+PAS+PPS",
        "internal_sensor_configuration": "IMS",
        "MCR_No": "MCR-2026-001",
        "role_summary": "PjM: Zhang San; TPM: Li Si; SW_PCM: Wang Wu",
        "role_email_summary": (
            "PjMEmail: Zhang San_zhang.san@example.com; "
            "TPMEmail: Li Si_li.si@example.com; "
            "SW_PCMEmail: Wang Wu_wang.wu@example.com"
        ),
    }
